from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

from app.services.docx_report import build_report_docx


def _image_bytes(color: str) -> bytes:
    image_stream = BytesIO()
    Image.new("RGB", (80, 60), color).save(image_stream, format="PNG")
    return image_stream.getvalue()


def test_building_model_detection_uses_model_template_and_nine_uploaded_images() -> None:
    overview_bytes = _image_bytes("red")
    report_image_bytes = _image_bytes("blue")
    images = [{
        "orientation": "overview",
        "image_kind": "model",
        "original_filename": "overview-model.png",
        "storage_bucket": "inspection",
        "storage_object_key": "overview-model.png",
    }, *[
        {
            "orientation": orientation,
            "image_kind": image_kind,
            "original_filename": f"{orientation}-{image_kind}.png",
            "storage_bucket": "inspection",
            "storage_object_key": f"{orientation}-{image_kind}.png",
        }
        for orientation in ("east", "west", "south", "north")
        for image_kind in ("elevation", "annotated")
    ]]

    content = build_report_docx(
        "测试报告",
        "RPT-MODEL-001",
        {
            "project": {"name": "测试建筑"},
            "detection_config": {
                "config_json": {"generate_building_model": True},
            },
            "building_model_images": images,
            "photos": [],
            "defects": [],
        },
        read_object=lambda _, object_key: (
            overview_bytes if object_key == "overview-model.png" else report_image_bytes
        ),
    )

    document = Document(BytesIO(content))
    assert document.paragraphs[0].text == "测试建筑无人机外立面表观病害筛查分析报告"
    assert len(document.tables) == 4
    assert len(document.inline_shapes) == 17
    overview_blip = document.paragraphs[3]._p.xpath(".//a:blip")[0]
    overview_part = document.part.related_parts[overview_blip.get(qn("r:embed"))]
    with Image.open(BytesIO(overview_part.blob)) as overview_image:
        assert overview_image.getpixel((0, 0)) == (255, 0, 0)


def test_building_model_result_tables_only_include_their_own_facade_photos() -> None:
    report_image_bytes = _image_bytes("blue")
    images = [{
        "orientation": "overview",
        "image_kind": "model",
        "original_filename": "overview-model.png",
        "storage_bucket": "inspection",
        "storage_object_key": "overview-model.png",
    }, *[
        {
            "orientation": orientation,
            "image_kind": image_kind,
            "original_filename": f"{orientation}-{image_kind}.png",
            "storage_bucket": "inspection",
            "storage_object_key": f"{orientation}-{image_kind}.png",
        }
        for orientation in ("east", "west", "south", "north")
        for image_kind in ("elevation", "annotated")
    ]]
    facade_names = ("东立面", "南立面", "西立面", "北立面")
    photos = [
        {
            "id": f"photo-{index}",
            "original_filename": f"facade-{index}_V.png",
            "photo_type": "visible",
            "facade_orientation": facade,
            "storage_bucket": "inspection",
            "storage_object_key": f"facade-{index}.png",
        }
        for index, facade in enumerate(facade_names, start=1)
    ]
    defects = [
        {
            "id": f"defect-{index}",
            "photo_id": photo["id"],
            "photo_filename": photo["original_filename"],
            "defect_type": "crack",
            "bbox_json": {
                "x": 100 * index,
                "y": 50 * index,
                "width": 80,
                "height": 40,
            },
            "length": 0.1 * index,
        }
        for index, photo in enumerate(photos, start=1)
    ]

    content = build_report_docx(
        "测试报告",
        "RPT-MODEL-002",
        {
            "project": {"name": "测试建筑"},
            "detection_config": {
                "config_json": {"generate_building_model": True},
            },
            "building_model_images": images,
            "photos": photos,
            "defects": defects,
        },
        read_object=lambda *_: report_image_bytes,
    )

    document = Document(BytesIO(content))
    assert len(document.tables) == 4
    for index, table in enumerate(document.tables, start=1):
        assert len(table.rows) == 2
        assert table.rows[1].cells[1].text.endswith(f"facade-{index}_V.png")
        assert "疑似裂缝: 1处" in table.rows[1].cells[3].text
        assert (
            f"（x={100 * index}，y={50 * index}）"
            in table.rows[1].cells[4].text
        )
